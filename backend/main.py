import base64
import io
import os
from typing import Any

import httpx
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from ai.provider import generate_text, get_ai_model, get_ai_provider
from ai.test_case_generator import generate_test_cases
from templates.analyzer import analyze_template

load_dotenv()

app = FastAPI(title="AI Test Case Generator API", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000", "https://ai-testing-tools.onrender.com"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DEFAULT_TEMPLATE = {
    "filename": "standard-qa",
    "format": "standard",
    "columns": [
        "TC ID", "Requirement ID", "Test Scenario", "Preconditions",
        "Test Steps", "Expected Result", "Priority", "Test Type",
    ],
}

class GenerateRequest(BaseModel):
    requirement: str = Field(..., min_length=1)
    template: dict[str, Any] = Field(default_factory=dict)
    # Optional context from additional requirements, acceptance criteria,
    # business rules, UI/UX evidence and API specifications.
    source_context: str = ""

class TextRequest(BaseModel):
    prompt: str = Field(..., min_length=1)

class DevOpsConfig(BaseModel):
    organization: str = Field(..., min_length=1)
    project: str = Field(..., min_length=1)
    personal_access_token: str = Field(..., min_length=1)

class DevOpsCreateRequest(BaseModel):
    config: DevOpsConfig
    test_cases: list[dict[str, Any]] = Field(default_factory=list)
    module_name: str = ""
    user_story_id: str = ""


def normalize_template(template: dict[str, Any] | None) -> dict[str, Any]:
    template = template if isinstance(template, dict) else {}
    columns = template.get("columns", [])
    if not isinstance(columns, list):
        columns = []
    columns = [str(x).strip() for x in columns if x is not None and str(x).strip()]
    if not columns:
        return {**DEFAULT_TEMPLATE, "columns": list(DEFAULT_TEMPLATE["columns"]), "column_count": len(DEFAULT_TEMPLATE["columns"])}
    return {**template, "columns": columns, "column_count": len(columns)}


def field_value(case: dict[str, Any], names: tuple[str, ...]) -> str:
    for key, value in case.items():
        if str(key).strip().lower() in {n.lower() for n in names} and value not in (None, ""):
            return str(value)
    return ""


def steps_to_html(steps: Any) -> str:
    if isinstance(steps, list):
        rows = []
        for i, step in enumerate(steps, 1):
            if isinstance(step, dict):
                action = step.get("action") or step.get("step") or step.get("description") or ""
                expected = step.get("expected") or step.get("expected_result") or ""
            else:
                action, expected = str(step), ""
            rows.append(f'<step id="{i}"><parameterizedString isformatted="true">{action}</parameterizedString><parameterizedString isformatted="true">{expected}</parameterizedString></step>')
        return "<steps id=\"0\" last=\"%d\">%s</steps>" % (len(rows), "".join(rows))
    return str(steps or "")


@app.get("/")
def root():
    return {"message": "AI Test Case Generator API", "status": "running", "provider": get_ai_provider(), "model": get_ai_model()}

@app.get("/api/health")
def health():
    return {"success": True, "provider": get_ai_provider(), "model": get_ai_model()}

@app.get("/api/ai/config")
def ai_config():
    return {"success": True, "provider": get_ai_provider(), "model": get_ai_model()}

@app.post("/api/ai/generate-text")
async def generate_text_endpoint(request: TextRequest):
    try:
        return {"success": True, "response": await generate_text(request.prompt.strip()), "provider": get_ai_provider(), "model": get_ai_model()}
    except Exception as exc:
        return {"success": False, "error": str(exc)}

@app.post("/api/ai/generate")
async def generate(request: GenerateRequest):
    requirement = request.requirement.strip()
    if not requirement:
        return {"success": False, "error": "Requirement is required."}
    template = normalize_template(request.template)

    # Keep the source hierarchy explicit at the backend boundary so the
    # generator receives the same QA rules even if another client calls the API.
    generation_input = f"""SOURCE AUTHORITY AND COVERAGE RULES

1. Explicit requirements, acceptance criteria and business rules supplied by the user are authoritative.
2. Supporting requirements/design/API material may clarify or expand coverage only where it is relevant to the supplied feature.
3. The QA knowledge library is guidance for selecting applicable testing dimensions. It must never introduce an unrelated requirement.
4. Never silently invent permissions, performance targets, security policies, workflows, validations, limits or expected outcomes that are not supported by the supplied sources.
5. When information is missing or ambiguous, flag it for QA input/ambiguity review rather than inventing a requirement.
6. Analyze applicable features, roles, workflows, validations, business rules, error scenarios, permissions, UI/UX behavior, API behavior, performance-sensitive operations and security risks.
7. Generate deep positive, negative, boundary, validation and workflow/state coverage where supported.
8. Every test case must be traceable to a supplied requirement, behavior, rule, acceptance criterion, UI/API behavior, or clearly derived behavior.
9. Use the supplied template as the source of truth for output columns, terminology and format.
10. Generate Priority during test-case creation when the selected template contains a Priority field.

SUPPLIED PRIMARY SOURCE MATERIAL

{requirement}

SUPPLIED SUPPORTING / DESIGN / API SOURCE MATERIAL

{request.source_context.strip() if request.source_context.strip() else "No additional supporting source was supplied."}
"""

    try:
        result = await generate_test_cases(generation_input, template)
        if not isinstance(result, dict):
            result = {}
        return {
            "success": True,
            "mode": "review",
            "saved": False,
            "provider": get_ai_provider(),
            "model": get_ai_model(),
            "template": template,
            "assumptions": result.get("assumptions", []),
            "coverage_analysis": result.get("coverage_analysis", {}),
            "test_cases": result.get("test_cases", []),
            "source_analysis": {
                "primary_requirement": True,
                "additional_supporting_inputs": bool(request.source_context.strip()),
                "supporting_source_context_included": bool(request.source_context.strip()),
                "knowledge_library": True,
                "rules": [
                    "Explicit supplied requirements are authoritative.",
                    "Knowledge-library guidance is applicability-bound and cannot create unrelated requirements.",
                    "Ambiguous or missing information is surfaced for QA input rather than silently invented.",
                ],
            },
        }
    except Exception as exc:
        print("Generation error:", repr(exc))
        return {"success": False, "error": str(exc)}

@app.post("/api/templates/analyze")
async def analyze_template_endpoint(file: UploadFile = File(...)):
    if not file.filename:
        return {"success": False, "error": "No template file selected."}
    try:
        content = await file.read()
        template = normalize_template(analyze_template(file.filename, content))
        return {"success": True, "template": template}
    except Exception as exc:
        return {"success": False, "error": f"Template analysis failed: {exc}"}

@app.post("/api/requirements/extract")
async def extract_requirement(file: UploadFile = File(...)):
    if not file.filename:
        return {"success": False, "error": "No requirement file selected."}
    filename = file.filename.lower()
    content = await file.read()
    if not content:
        return {"success": False, "error": "The uploaded file is empty."}
    try:
        text = _extract_document_text(file.filename, content).strip()
        if not text:
            return {
                "success": False,
                "error": (
                    "Supported requirement files: TXT, MD, CSV, XLSX, DOCX or PDF. "
                    "No readable text was found."
                ),
            }
        return {"success": True, "filename": file.filename, "text": text}
    except Exception as exc:
        return {"success": False, "error": f"Requirement extraction failed: {exc}"}

def _extract_document_text(filename: str, content: bytes) -> str:
    """Extract readable text from supported requirement/supporting files."""
    lower = filename.lower()

    if lower.endswith((".txt", ".md", ".markdown", ".csv")):
        return content.decode("utf-8-sig")

    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v not in (None, "")]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    if lower.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Include table content because business rules/acceptance criteria are
        # frequently stored in Word tables.
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    paragraphs.append(" | ".join(values))

        return "\n".join(paragraphs)

    if lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    return ""


async def _analyze_image_with_gemini(filename: str, content: bytes) -> str:
    """Describe a supplied UI/UX screenshot so it can participate in QA analysis."""
    provider = get_ai_provider()

    if provider != "gemini":
        return (
            f"UI/UX screenshot supplied: {filename}. "
            "Visual analysis is unavailable for the currently selected AI provider."
        )

    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    model = get_ai_model()

    if not api_key:
        return (
            f"UI/UX screenshot supplied: {filename}. "
            "GEMINI_API_KEY is not configured, so visual analysis could not be performed."
        )

    try:
        from google import genai
        from google.genai import types

        mime_map = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".webp": "image/webp",
        }
        mime_type = mime_map.get(
            os.path.splitext(filename.lower())[1],
            "image/png",
        )

        client = genai.Client(api_key=api_key)
        prompt = f"""
Analyze this UI/UX screenshot only as evidence for QA test generation.

Filename: {filename}

Extract only observable, test-relevant information:
- visible screens/pages and major sections
- labels, buttons, links, fields, dropdowns, checkboxes and controls
- visible validation/error/help states
- navigation/workflow clues
- visible roles/permissions clues
- UI states and state transitions
- important layout/responsive behavior that is explicitly observable
- any text that appears to define expected behavior

Do not invent hidden functionality, backend rules, permissions, performance targets,
security requirements or business rules that cannot be supported by the screenshot.
Clearly say when something is not observable.
Return concise structured QA evidence.
"""
        response = await client.aio.models.generate_content(
            model=model,
            contents=[
                types.Part.from_bytes(data=content, mime_type=mime_type),
                prompt,
            ],
        )
        return str(response.text or "").strip()
    except Exception as exc:
        return (
            f"UI/UX screenshot supplied: {filename}. "
            f"Visual analysis failed safely: {exc}"
        )


@app.post("/api/requirements/extract-many")
async def extract_supporting_inputs(
    files: list[UploadFile] = File(...),
):
    """
    Extract multiple optional supporting QA inputs in one request.

    Documents are converted to text. UI/UX images are analyzed with Gemini
    when Gemini is the active provider. Nothing is persisted.
    """
    if not files:
        return {"success": False, "error": "No supporting files selected."}

    allowed_document_extensions = {
        ".txt", ".md", ".markdown", ".csv", ".xlsx", ".docx", ".pdf"
    }
    allowed_image_extensions = {".png", ".jpg", ".jpeg", ".webp"}

    extracted: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    for file in files:
        filename = (file.filename or "").strip()
        if not filename:
            errors.append({"filename": "", "error": "A selected file has no filename."})
            continue

        content = await file.read()
        if not content:
            errors.append({"filename": filename, "error": "The uploaded file is empty."})
            continue

        extension = os.path.splitext(filename.lower())[1]

        try:
            if extension in allowed_document_extensions:
                text = _extract_document_text(filename, content).strip()
                if not text:
                    errors.append({
                        "filename": filename,
                        "error": "No readable text was found in the uploaded file.",
                    })
                    continue

                extracted.append({
                    "filename": filename,
                    "type": "document",
                    "text": text,
                })

            elif extension in allowed_image_extensions:
                text = await _analyze_image_with_gemini(filename, content)
                extracted.append({
                    "filename": filename,
                    "type": "ui_ux_screenshot",
                    "text": text,
                })

            else:
                errors.append({
                    "filename": filename,
                    "error": (
                        "Supported files: TXT, MD, CSV, XLSX, DOCX, PDF, "
                        "PNG, JPG or WEBP."
                    ),
                })
        except Exception as exc:
            errors.append({
                "filename": filename,
                "error": f"Extraction failed: {exc}",
            })

    if not extracted:
        return {
            "success": False,
            "error": "No supporting input could be read.",
            "files": [],
            "errors": errors,
        }

    return {
        "success": True,
        "files": extracted,
        "errors": errors,
    }


def ado_url(config: DevOpsConfig, path: str) -> str:
    org = config.organization.strip().strip("/")
    project = config.project.strip().strip("/")

    if path.startswith("_apis/projects"):
        return f"https://dev.azure.com/{org}/{path}"

    return f"https://dev.azure.com/{org}/{project}/{path}"

async def ado_request(config: DevOpsConfig, method: str, url: str, **kwargs):
    token = base64.b64encode(f":{config.personal_access_token}".encode()).decode()
    headers = kwargs.pop("headers", {})
    headers.update({"Authorization": f"Basic {token}", "Content-Type": "application/json-patch+json"})
    async with httpx.AsyncClient(timeout=45) as client:
        response = await client.request(method, url, headers=headers, **kwargs)
        if response.status_code >= 400:
            raise RuntimeError(f"Azure DevOps returned {response.status_code}: {response.text[:500]}")
        return response.json() if response.content else {}

@app.post("/api/devops/test-connection")
async def devops_test_connection(config: DevOpsConfig):
    try:
        url = ado_url(config, "_apis/projects?api-version=7.1")
        data = await ado_request(config, "GET", url, headers={"Content-Type": "application/json"})
        return {"success": True, "message": "Azure DevOps connection successful.", "project_count": len(data.get("value", []))}
    except Exception as exc:
        return {"success": False, "error": str(exc)}

@app.post("/api/devops/create-test-cases")
async def devops_create_test_cases(request: DevOpsCreateRequest):
    if not request.test_cases:
        return {"success": False, "error": "No test cases to send."}
    created = []
    failed = []
    for case in request.test_cases:
        title = field_value(case, ("title", "test scenario", "scenario", "test case", "business check")) or "AI Generated Test Case"
        steps = field_value(case, ("test steps", "steps", "test step"))
        expected = field_value(case, ("expected result", "expected outcome", "result"))
        module = request.module_name.strip()
        if module:
            title = f"[{module}] {title}"

        # Avoid creating an obvious duplicate already present in Azure DevOps.
        try:
            escaped_title = title.replace("'", "''")
            wiql = {
                "query": (
                    "SELECT [System.Id] FROM WorkItems "
                    "WHERE [System.TeamProject] = @project "
                    "AND [System.WorkItemType] = 'Test Case' "
                    f"AND [System.Title] = '{escaped_title}'"
                )
            }
            wiql_url = ado_url(request.config, "_apis/wit/wiql?api-version=7.1")
            token = base64.b64encode(f":{request.config.personal_access_token}".encode()).decode()
            async with httpx.AsyncClient(timeout=45) as client:
                check = await client.post(
                    wiql_url,
                    headers={"Authorization": f"Basic {token}", "Content-Type": "application/json"},
                    json=wiql,
                )
            if check.status_code < 400 and check.json().get("workItems"):
                created.append({"id": check.json()["workItems"][0].get("id"), "title": title, "skipped_duplicate": True})
                continue
        except Exception:
            # Publishing should continue if the duplicate lookup is unavailable.
            pass

        patch = [
            {"op": "add", "path": "/fields/System.Title", "value": title},
            {"op": "add", "path": "/fields/Microsoft.VSTS.TCM.Steps", "value": steps_to_html(steps or expected)},
        ]
        try:
            url = ado_url(request.config, "_apis/wit/workitems/$Test%20Case?api-version=7.1")
            item = await ado_request(request.config, "POST", url, json=patch)
            work_item_id = item.get("id")
            if request.user_story_id.strip() and work_item_id:
                link_patch = [{"op": "add", "path": "/relations/-", "value": {"rel": "Microsoft.VSTS.Common.TestedBy-Reverse", "url": item.get("url")}}]
                try:
                    await ado_request(request.config, "PATCH", f"{ado_url(request.config, f'_apis/wit/workitems/{request.user_story_id.strip()}?api-version=7.1')}", json=link_patch)
                except Exception:
                    pass
            created.append({"id": work_item_id, "title": title})
        except Exception as exc:
            failed.append({"title": title, "error": str(exc)})
    return {"success": len(created) > 0 and not failed, "created": created, "failed": failed, "created_count": len(created), "failed_count": len(failed)}

