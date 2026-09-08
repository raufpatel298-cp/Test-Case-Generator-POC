import csv
import io
import re
from typing import Any

from openpyxl import load_workbook


def analyze_template(
    filename: str,
    content: bytes,
) -> dict[str, Any]:

    if not filename:
        raise ValueError(
            "Template filename is required."
        )

    lower_filename = filename.lower()

    if lower_filename.endswith(".xlsx"):
        result = _analyze_xlsx(content)

    elif lower_filename.endswith(".csv"):
        result = _analyze_csv(content)

    elif lower_filename.endswith(".xls"):
        raise ValueError(
            "Legacy .xls files are not supported yet. "
            "Please save the template as .xlsx."
        )

    else:
        raise ValueError(
            "Unsupported template format. "
            "Please upload .xlsx or .csv."
        )

    result["filename"] = filename

    return result


# ==================================================
# XLSX
# ==================================================

def _analyze_xlsx(
    content: bytes,
) -> dict[str, Any]:

    try:
        workbook = load_workbook(
            filename=io.BytesIO(content),
            read_only=False,
            data_only=False,
        )
    except Exception as error:
        raise ValueError(
            "Could not read the Excel template."
        ) from error

    try:

        if not workbook.sheetnames:
            raise ValueError(
                "The Excel template contains no worksheets."
            )

        sheets = []

        for worksheet in workbook.worksheets:

            # --------------------------------------
            # Find header row
            # --------------------------------------

            header_row_number = _find_header_row(
                worksheet
            )

            header_values = [
                cell.value
                for cell in worksheet[
                    header_row_number
                ]
            ]

            columns = _clean_headers(
                header_values
            )

            if not columns:
                continue

            # --------------------------------------
            # Analyze columns
            # --------------------------------------

            column_metadata = []

            for index, cell in enumerate(
                worksheet[
                    header_row_number
                ],
                start=1,
            ):

                if index > len(columns):
                    break

                column_name = columns[
                    index - 1
                ]

                data_validation_values = (
                    _get_data_validation_values(
                        worksheet,
                        workbook,
                        cell.column_letter,
                        header_row_number,
                    )
                )

                column_metadata.append(
                    {
                        "index": index,
                        "column": column_name,
                        "excel_column": (
                            cell.column_letter
                        ),
                        "header": (
                            ""
                            if cell.value is None
                            else str(cell.value)
                        ),
                        "required": (
                            _infer_required_field(
                                worksheet,
                                cell.column_letter,
                                header_row_number,
                            )
                        ),
                        "data_type": (
                            _infer_column_data_type(
                                worksheet,
                                index,
                                header_row_number,
                            )
                        ),
                        "sample_values": (
                            _get_column_sample_values(
                                worksheet,
                                index,
                                header_row_number,
                            )
                        ),
                        "allowed_values": (
                            data_validation_values
                        ),
                        "number_format": (
                            cell.number_format
                        ),
                        "width": (
                            worksheet.column_dimensions[
                                cell.column_letter
                            ].width
                        ),
                    }
                )

            # --------------------------------------
            # Existing template data
            # --------------------------------------

            sample_rows = []

            max_sample_rows = 5

            for row in worksheet.iter_rows(
                min_row=header_row_number + 1,
                values_only=True,
            ):

                values = list(row)

                if not any(
                    value is not None
                    and str(value).strip() != ""
                    for value in values
                ):
                    continue

                normalized_row = {}

                for index, column in enumerate(
                    columns
                ):

                    value = (
                        values[index]
                        if index < len(values)
                        else ""
                    )

                    normalized_row[
                        column
                    ] = value

                sample_rows.append(
                    normalized_row
                )

                if (
                    len(sample_rows)
                    >= max_sample_rows
                ):
                    break

            # --------------------------------------
            # Template structure
            # --------------------------------------

            field_definitions = (
                _build_field_definitions(
                    column_metadata
                )
            )

            test_step_structure = (
                _analyze_test_step_structure(
                    worksheet,
                    columns,
                    header_row_number,
                )
            )

            expected_result_structure = (
                _analyze_expected_result_structure(
                    worksheet,
                    columns,
                    header_row_number,
                )
            )

            categorization_structure = (
                _analyze_categorization_structure(
                    column_metadata
                )
            )

            formatting_rules = (
                _analyze_formatting(
                    worksheet,
                    columns,
                    header_row_number,
                )
            )

            # --------------------------------------
            # Sheet metadata
            # --------------------------------------

            sheets.append(
                {
                    "name": worksheet.title,
                    "columns": columns,
                    "column_count": len(columns),
                    "header_row": (
                        header_row_number
                    ),
                    "column_metadata": (
                        column_metadata
                    ),
                    "sample_rows": sample_rows,
                    "field_definitions": (
                        field_definitions
                    ),
                    "test_step_structure": (
                        test_step_structure
                    ),
                    "expected_result_structure": (
                        expected_result_structure
                    ),
                    "categorization_structure": (
                        categorization_structure
                    ),
                    "formatting_rules": (
                        formatting_rules
                    ),
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                }
            )

        if not sheets:
            raise ValueError(
                "The Excel template contains no usable "
                "header columns."
            )

        first_sheet = sheets[0]

        return {
            "format": "xlsx",

            "sheet_names": workbook.sheetnames,

            "sheets": sheets,

            "columns": first_sheet[
                "columns"
            ],

            "column_count": first_sheet[
                "column_count"
            ],

            "header_row": first_sheet[
                "header_row"
            ],

            "column_metadata": first_sheet[
                "column_metadata"
            ],

            "sample_rows": first_sheet[
                "sample_rows"
            ],

            "field_definitions": first_sheet[
                "field_definitions"
            ],

            "test_step_structure": first_sheet[
                "test_step_structure"
            ],

            "expected_result_structure": first_sheet[
                "expected_result_structure"
            ],

            "categorization_structure": first_sheet[
                "categorization_structure"
            ],

            "formatting_rules": first_sheet[
                "formatting_rules"
            ],
        }

    finally:
        workbook.close()


# ==================================================
# Header Row
# ==================================================

def _find_header_row(
    worksheet,
) -> int:

    max_scan_rows = min(
        worksheet.max_row,
        10,
    )

    best_row = 1
    best_score = -1

    for row_number in range(
        1,
        max_scan_rows + 1,
    ):

        values = [
            cell.value
            for cell in worksheet[
                row_number
            ]
        ]

        non_empty = [
            value
            for value in values
            if value is not None
            and str(value).strip() != ""
        ]

        if not non_empty:
            continue

        # Prefer rows containing common
        # test-case terminology.

        text = " ".join(
            str(value).strip().lower()
            for value in non_empty
        )

        keyword_score = sum(
            1
            for keyword in (
                "test",
                "case",
                "scenario",
                "step",
                "expected",
                "result",
                "requirement",
                "priority",
                "severity",
                "status",
            )
            if keyword in text
        )

        score = (
            len(non_empty) * 10
            + keyword_score * 5
        )

        if score > best_score:
            best_score = score
            best_row = row_number

    return best_row


# ==================================================
# CSV
# ==================================================

def _analyze_csv(
    content: bytes,
) -> dict[str, Any]:

    try:
        text = content.decode(
            "utf-8-sig"
        )
    except UnicodeDecodeError as error:
        raise ValueError(
            "CSV template must be UTF-8 encoded."
        ) from error

    if not text.strip():
        raise ValueError(
            "The CSV template is empty."
        )

    reader = csv.reader(
        io.StringIO(text)
    )

    try:
        header_row = next(reader)
    except StopIteration:
        raise ValueError(
            "The CSV template contains no rows."
        )

    headers = _clean_headers(
        header_row
    )

    if not headers:
        raise ValueError(
            "The CSV template contains no usable columns."
        )

    # ----------------------------------------------
    # Capture sample rows
    # ----------------------------------------------

    sample_rows = []

    for row in reader:

        if not any(
            value is not None
            and str(value).strip() != ""
            for value in row
        ):
            continue

        normalized_row = {}

        for index, column in enumerate(
            headers
        ):

            value = (
                row[index]
                if index < len(row)
                else ""
            )

            normalized_row[
                column
            ] = value

        sample_rows.append(
            normalized_row
        )

        if len(sample_rows) >= 5:
            break

    # ----------------------------------------------
    # Column metadata
    # ----------------------------------------------

    column_metadata = []

    for index, column in enumerate(
        headers,
        start=1,
    ):

        sample_values = []

        for row in sample_rows:

            value = row.get(
                column,
                "",
            )

            if (
                value is not None
                and str(value).strip()
            ):

                sample_values.append(
                    value
                )

        column_metadata.append(
            {
                "index": index,
                "column": column,
                "excel_column": None,
                "header": column,
                "required": False,
                "data_type": (
                    _infer_values_data_type(
                        sample_values
                    )
                ),
                "sample_values": (
                    sample_values
                ),
                "allowed_values": (
                    _detect_allowed_values(
                        sample_values
                    )
                ),
                "number_format": None,
                "width": None,
            }
        )

    field_definitions = (
        _build_field_definitions(
            column_metadata
        )
    )

    test_step_structure = (
        _infer_structure_from_column_names(
            headers,
            (
                "test_step",
                "test steps",
                "test step",
                "steps",
                "step",
            ),
        )
    )

    expected_result_structure = (
        _infer_structure_from_column_names(
            headers,
            (
                "expected result",
                "expected_result",
                "expected",
                "result",
            ),
        )
    )

    categorization_structure = (
        _analyze_categorization_structure(
            column_metadata
        )
    )

    return {
        "format": "csv",

        "sheet_names": [],

        "sheets": [],

        "columns": headers,

        "column_count": len(headers),

        "header_row": 1,

        "column_metadata": column_metadata,

        "sample_rows": sample_rows,

        "field_definitions": (
            field_definitions
        ),

        "test_step_structure": (
            test_step_structure
        ),

        "expected_result_structure": (
            expected_result_structure
        ),

        "categorization_structure": (
            categorization_structure
        ),

        "formatting_rules": {
            "source": "csv",
            "formatting_available": False,
        },
    }


# ==================================================
# Required Field Detection
# ==================================================

def _infer_required_field(
    worksheet,
    column_letter: str,
    header_row: int,
) -> bool:

    # Look at existing data rows.
    #
    # This is intentionally conservative.
    # A populated column does not automatically mean
    # that the organization requires the field.

    values = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(
            worksheet.max_row,
            header_row + 20,
        ),
        min_col=worksheet[
            f"{column_letter}{header_row}"
        ].column,
        max_col=worksheet[
            f"{column_letter}{header_row}"
        ].column,
        values_only=True,
    ):

        if row:
            values.append(
                row[0]
            )

    non_empty = [
        value
        for value in values
        if value is not None
        and str(value).strip() != ""
    ]

    if not values:
        return False

    # Only infer required when all observed
    # example rows contain a value.
    return len(non_empty) == len(values)


# ==================================================
# Data Type Detection
# ==================================================

def _infer_column_data_type(
    worksheet,
    column_index: int,
    header_row: int,
) -> str:

    values = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(
            worksheet.max_row,
            header_row + 20,
        ),
        min_col=column_index,
        max_col=column_index,
        values_only=True,
    ):

        if row:
            value = row[0]

            if value is not None:
                values.append(value)

    return _infer_values_data_type(
        values
    )


def _infer_values_data_type(
    values,
) -> str:

    if not values:
        return "unknown"

    if all(
        isinstance(
            value,
            bool,
        )
        for value in values
    ):
        return "boolean"

    if all(
        isinstance(
            value,
            (int, float),
        )
        and not isinstance(
            value,
            bool,
        )
        for value in values
    ):
        return "number"

    return "text"


# ==================================================
# Sample Values
# ==================================================

def _get_column_sample_values(
    worksheet,
    column_index: int,
    header_row: int,
) -> list[Any]:

    values = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(
            worksheet.max_row,
            header_row + 20,
        ),
        min_col=column_index,
        max_col=column_index,
        values_only=True,
    ):

        if not row:
            continue

        value = row[0]

        if (
            value is not None
            and str(value).strip() != ""
        ):

            if value not in values:
                values.append(
                    value
                )

        if len(values) >= 10:
            break

    return values


# ==================================================
# Data Validation
# ==================================================

def _get_data_validation_values(
    worksheet,
    workbook,
    column_letter: str,
    header_row: int,
) -> list[str]:
    """
    Extract explicit Excel data-validation list
    values for a template column.

    Supported forms include:

    1. Inline lists:
       P1,P2,P3

    2. Quoted inline lists:
       "P1,P2,P3"

    3. Worksheet cell ranges:
       =Lists!$A$1:$A$3

    4. Workbook defined names that resolve to
       worksheet ranges.
    """

    values = []

    try:

        for validation in (
            worksheet.data_validations.dataValidation
        ):

            formula = validation.formula1

            if not formula:
                continue

            if not _validation_applies_to_column(
                validation,
                column_letter,
            ):
                continue

            formula_text = str(
                formula
            ).strip()

            # ------------------------------------------
            # Inline list
            # ------------------------------------------

            inline_values = (
                _parse_inline_validation_list(
                    formula_text
                )
            )

            if inline_values:
                values.extend(
                    inline_values
                )
                continue

            # ------------------------------------------
            # Formula / range / named range
            # ------------------------------------------

            formula_values = (
                _resolve_validation_formula(
                    formula_text,
                    workbook,
                    worksheet,
                )
            )

            if formula_values:
                values.extend(
                    formula_values
                )

    except Exception:
        # Template analysis should not fail simply
        # because a data-validation rule is unusual.
        pass

    # Preserve order while removing duplicates.

    unique_values = []

    for value in values:

        cleaned_value = str(
            value
        ).strip()

        if (
            cleaned_value
            and cleaned_value not in unique_values
        ):

            unique_values.append(
                cleaned_value
            )

    return unique_values


# ==================================================
# Validation Range Matching
# ==================================================

def _validation_applies_to_column(
    validation,
    column_letter: str,
) -> bool:

    ranges = str(
        validation.sqref
    )

    if not ranges:
        return False

    for reference in ranges.split():

        reference = reference.strip()

        if not reference:
            continue

        # Extract Excel column letters from
        # a validation range such as:
        # G2:G100
        #
        # Also handles:
        # $G$2:$G$100

        columns = re.findall(
            r"\$?([A-Z]{1,3})\$?\d+",
            reference.upper(),
        )

        if not columns:
            continue

        if column_letter.upper() in columns:
            return True

    return False


# ==================================================
# Inline Validation List
# ==================================================

def _parse_inline_validation_list(
    formula_text: str,
) -> list[str]:

    value = formula_text.strip()

    if not value:
        return []

    # Excel may represent inline validation as:
    #
    # "P1,P2,P3"
    #
    # or the formula may arrive without quotes:
    #
    # P1,P2,P3

    if (
        value.startswith('"')
        and value.endswith('"')
    ):

        value = value[1:-1]

    # Ignore formulas that clearly reference
    # cells, ranges, or functions.

    if (
        value.startswith("=")
        or "!" in value
        or "$" in value
        or "(" in value
        or ")" in value
    ):
        return []

    if "," not in value:
        return []

    return [
        item.strip()
        for item in value.split(",")
        if item.strip()
    ]


# ==================================================
# Validation Formula Resolution
# ==================================================

def _resolve_validation_formula(
    formula_text: str,
    workbook,
    worksheet,
) -> list[str]:

    formula = formula_text.strip()

    if not formula:
        return []

    if formula.startswith("="):
        formula = formula[1:].strip()

    # ----------------------------------------------
    # Worksheet range
    # ----------------------------------------------

    range_match = re.match(
        r"^(?:'([^']+)'|([A-Za-z0-9_ .-]+))!"
        r"\$?([A-Z]{1,3})\$?(\d+)"
        r":\$?([A-Z]{1,3})\$?(\d+)$",
        formula,
    )

    if range_match:

        quoted_sheet = range_match.group(
            1
        )

        plain_sheet = range_match.group(
            2
        )

        sheet_name = (
            quoted_sheet
            if quoted_sheet is not None
            else plain_sheet
        )

        start_column = range_match.group(
            3
        )

        start_row = int(
            range_match.group(4)
        )

        end_column = range_match.group(
            5
        )

        end_row = int(
            range_match.group(6)
        )

        if sheet_name not in workbook.sheetnames:
            return []

        source_sheet = workbook[
            sheet_name
        ]

        values = []

        for row in source_sheet.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=_excel_column_number(
                start_column
            ),
            max_col=_excel_column_number(
                end_column
            ),
            values_only=True,
        ):

            for value in row:

                if (
                    value is not None
                    and str(value).strip()
                ):

                    values.append(
                        str(value).strip()
                    )

        return values

    # ----------------------------------------------
    # Same-sheet range without sheet name
    # ----------------------------------------------

    local_range_match = re.match(
        r"^\$?([A-Z]{1,3})\$?(\d+)"
        r":\$?([A-Z]{1,3})\$?(\d+)$",
        formula,
    )

    if local_range_match:

        start_column = (
            local_range_match.group(1)
        )

        start_row = int(
            local_range_match.group(2)
        )

        end_column = (
            local_range_match.group(3)
        )

        end_row = int(
            local_range_match.group(4)
        )

        values = []

        for row in worksheet.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=_excel_column_number(
                start_column
            ),
            max_col=_excel_column_number(
                end_column
            ),
            values_only=True,
        ):

            for value in row:

                if (
                    value is not None
                    and str(value).strip()
                ):

                    values.append(
                        str(value).strip()
                    )

        return values

    # ----------------------------------------------
    # Workbook defined name
    # ----------------------------------------------

    defined_name_values = (
        _resolve_defined_name(
            formula,
            workbook,
        )
    )

    if defined_name_values:
        return defined_name_values

    return []


# ==================================================
# Defined Name Resolution
# ==================================================

def _resolve_defined_name(
    name: str,
    workbook,
) -> list[str]:

    try:

        defined_name = (
            workbook.defined_names.get(
                name
            )
        )

        if defined_name is None:
            return []

        destinations = (
            list(
                defined_name.destinations
            )
        )

        values = []

        for sheet_name, cell_range in (
            destinations
        ):

            if sheet_name not in workbook.sheetnames:
                continue

            worksheet = workbook[
                sheet_name
            ]

            for row in worksheet[
                cell_range
            ]:

                for cell in row:

                    if (
                        cell.value is not None
                        and str(
                            cell.value
                        ).strip()
                    ):

                        values.append(
                            str(
                                cell.value
                            ).strip()
                        )

        return values

    except Exception:
        return []


# ==================================================
# Excel Column Number
# ==================================================

def _excel_column_number(
    column: str,
) -> int:

    number = 0

    for character in column.upper():

        if not (
            "A" <= character <= "Z"
        ):
            continue

        number = (
            number * 26
            + (
                ord(character)
                - ord("A")
                + 1
            )
        )

    return number


# ==================================================
# Allowed Value Detection
# ==================================================

def _detect_allowed_values(
    sample_values,
) -> list[str]:

    if not sample_values:
        return []

    unique_values = []

    for value in sample_values:

        text = str(value).strip()

        if (
            text
            and text not in unique_values
        ):

            unique_values.append(
                text
            )

    # Only expose a small finite set as
    # candidate allowed values.
    #
    # These are observations, not invented rules.

    if (
        len(unique_values) <= 10
    ):

        return unique_values

    return []


# ==================================================
# Field Definitions
# ==================================================

def _build_field_definitions(
    column_metadata,
) -> dict[str, Any]:

    definitions = {}

    for metadata in column_metadata:

        column = metadata.get(
            "column"
        )

        if not column:
            continue

        definitions[
            column
        ] = {
            "required": bool(
                metadata.get(
                    "required",
                    False,
                )
            ),
            "data_type": metadata.get(
                "data_type",
                "unknown",
            ),
            "sample_values": metadata.get(
                "sample_values",
                [],
            ),
            "allowed_values": metadata.get(
                "allowed_values",
                [],
            ),
        }

    return definitions


# ==================================================
# Test Step Structure
# ==================================================

def _analyze_test_step_structure(
    worksheet,
    columns,
    header_row: int,
) -> dict[str, Any]:

    step_columns = [
        column
        for column in columns
        if _is_step_column(column)
    ]

    if not step_columns:

        return {
            "present": False,
            "columns": [],
            "style": "unknown",
            "observed_samples": [],
        }

    samples = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(
            worksheet.max_row,
            header_row + 6,
        ),
        values_only=True,
    ):

        row_values = list(row)

        for column in step_columns:

            index = columns.index(
                column
            )

            if index >= len(row_values):
                continue

            value = row_values[index]

            if (
                value is None
                or not str(value).strip()
            ):
                continue

            samples.append(
                str(value)
            )

    style = "single_field"

    if any(
        "\n" in sample
        for sample in samples
    ):
        style = "multi_line"

    return {
        "present": True,
        "columns": step_columns,
        "style": style,
        "observed_samples": samples[:5],
    }


# ==================================================
# Expected Result Structure
# ==================================================

def _analyze_expected_result_structure(
    worksheet,
    columns,
    header_row: int,
) -> dict[str, Any]:

    result_columns = [
        column
        for column in columns
        if _is_expected_result_column(
            column
        )
    ]

    if not result_columns:

        return {
            "present": False,
            "columns": [],
            "style": "unknown",
            "observed_samples": [],
        }

    samples = []

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=min(
            worksheet.max_row,
            header_row + 6,
        ),
        values_only=True,
    ):

        row_values = list(row)

        for column in result_columns:

            index = columns.index(
                column
            )

            if index >= len(row_values):
                continue

            value = row_values[index]

            if (
                value is None
                or not str(value).strip()
            ):
                continue

            samples.append(
                str(value)
            )

    style = "single_field"

    if any(
        "\n" in sample
        for sample in samples
    ):
        style = "multi_line"

    return {
        "present": True,
        "columns": result_columns,
        "style": style,
        "observed_samples": samples[:5],
    }


# ==================================================
# Categorization
# ==================================================

def _analyze_categorization_structure(
    column_metadata,
) -> dict[str, Any]:

    categories = {}

    for metadata in column_metadata:

        column = str(
            metadata.get(
                "column",
                "",
            )
        )

        normalized = (
            column.lower()
            .strip()
        )

        if any(
            keyword in normalized
            for keyword in (
                "priority",
                "severity",
                "test type",
                "test_type",
                "category",
                "status",
                "classification",
            )
        ):

            categories[
                column
            ] = {
                "sample_values": (
                    metadata.get(
                        "sample_values",
                        [],
                    )
                ),
                "allowed_values": (
                    metadata.get(
                        "allowed_values",
                        [],
                    )
                ),
            }

    return {
        "fields": categories
    }


# ==================================================
# Formatting
# ==================================================

def _analyze_formatting(
    worksheet,
    columns,
    header_row: int,
) -> dict[str, Any]:

    header_cells = list(
        worksheet[
            header_row
        ]
    )

    widths = {}

    for cell in header_cells:

        if cell.column > len(columns):
            continue

        column = columns[
            cell.column - 1
        ]

        widths[
            column
        ] = worksheet.column_dimensions[
            cell.column_letter
        ].width

    merged_ranges = [
        str(item)
        for item in worksheet.merged_cells.ranges
    ]

    freeze_panes = (
        str(worksheet.freeze_panes)
        if worksheet.freeze_panes
        else None
    )

    return {
        "sheet_name": worksheet.title,
        "header_row": header_row,
        "column_widths": widths,
        "merged_ranges": merged_ranges,
        "freeze_panes": freeze_panes,
        "header_style": {
            "bold": any(
                cell.font.bold
                for cell in header_cells
                if cell.value is not None
            ),
            "has_fill": any(
                cell.fill.fill_type
                for cell in header_cells
                if cell.value is not None
            ),
        },
    }


# ==================================================
# Structure Helpers
# ==================================================

def _is_step_column(
    column: str,
) -> bool:

    normalized = (
        str(column)
        .strip()
        .lower()
    )

    return normalized in {
        "step",
        "steps",
        "test step",
        "test steps",
        "test_step",
        "test_steps",
        "action",
        "actions",
        "procedure",
        "procedure steps",
    }


def _is_expected_result_column(
    column: str,
) -> bool:

    normalized = (
        str(column)
        .strip()
        .lower()
    )

    return normalized in {
        "expected result",
        "expected results",
        "expected_result",
        "expected_results",
        "expected",
        "result",
        "expected behavior",
        "expected behaviour",
    }


def _infer_structure_from_column_names(
    columns,
    candidates,
) -> dict[str, Any]:

    matches = []

    normalized_candidates = {
        str(value).strip().lower()
        for value in candidates
    }

    for column in columns:

        normalized = (
            str(column)
            .strip()
            .lower()
        )

        if normalized in normalized_candidates:
            matches.append(
                column
            )

    return {
        "present": bool(matches),
        "columns": matches,
        "style": (
            "single_field"
            if matches
            else "unknown"
        ),
    }


# ==================================================
# Header Cleaning
# ==================================================

def _clean_headers(
    headers,
) -> list[str]:

    cleaned = []

    for index, header in enumerate(
        headers
    ):

        value = (
            ""
            if header is None
            else str(header).strip()
        )

        if not value:
            value = (
                f"Unnamed Column {index + 1}"
            )

        # Preserve original spelling for
        # unique columns.
        #
        # Duplicate headers are made unique
        # only because JSON objects cannot safely
        # represent duplicate keys.

        original_value = value
        duplicate_count = 1

        while value in cleaned:

            duplicate_count += 1

            value = (
                f"{original_value} "
                f"({duplicate_count})"
            )

        cleaned.append(value)

    return cleaned